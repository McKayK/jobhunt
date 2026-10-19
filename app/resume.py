"""Resume files in and out.

In:  PDF / DOCX / TXT / MD -> plain text (what matching and tailoring use).
Out: a tailored resume (the JSON shape from ai.tailor) -> a clean one-page
     .docx in the same layout as a typical developer resume, plus a plain-text
     version for pasting into application forms.
"""
from __future__ import annotations

import io
import re
from typing import Any


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    elif name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(lines)
    else:
        text = data.decode("utf-8", errors="replace")
    text = text.replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --- Plain text -------------------------------------------------------------

def to_text(r: dict) -> str:
    out: list[str] = []
    if r.get("name"):
        out.append(r["name"])
    if r.get("headline"):
        out.append(r["headline"])
    for c in r.get("contact") or []:
        out.append(c)
    if r.get("summary"):
        out += ["", "SUMMARY", r["summary"]]
    if r.get("skills"):
        out += ["", "TECHNICAL SKILLS"]
        for s in r["skills"]:
            items = ", ".join(s.get("items") or [])
            out.append(f"{s.get('category')}: {items}" if s.get("category") else items)
    if r.get("experience"):
        out += ["", "EXPERIENCE"]
        for e in r["experience"]:
            head = " | ".join(x for x in [e.get("company"), e.get("location")] if x)
            out.append(f"{head}    {e.get('dates') or ''}".rstrip())
            if e.get("title"):
                out.append(e["title"])
            for b in e.get("bullets") or []:
                out.append(f"  • {b}")
            for p in e.get("projects") or []:
                out.append(" | ".join(x for x in [p.get("name"), p.get("tech")] if x))
                for b in p.get("bullets") or []:
                    out.append(f"  • {b}")
    if r.get("projects"):
        out += ["", "PROJECTS"]
        for p in r["projects"]:
            out.append(" | ".join(x for x in [p.get("name"), p.get("tech")] if x))
            for b in p.get("bullets") or []:
                out.append(f"  • {b}")
    if r.get("education"):
        out += ["", "EDUCATION"] + list(r["education"])
    return "\n".join(out).strip() + "\n"


# --- DOCX -------------------------------------------------------------------

def to_docx(r: dict) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.6))
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    text_width = sec.page_width - sec.left_margin - sec.right_margin

    base = d.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10)
    base.paragraph_format.space_after = Pt(0)
    base.paragraph_format.space_before = Pt(0)

    def para(text: str = "", bold=False, size=None, align=None, italic=False,
             space_before=0, space_after=0, color=None):
        p = d.add_paragraph()
        if text:
            run = p.add_run(text)
            run.bold, run.italic = bold, italic
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor(*color)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def heading(text: str):
        p = para(text.upper(), bold=True, size=10.5, space_before=7, space_after=2)
        # Bottom border under section headings.
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "444444")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)

    def left_right(left: str, right: str, bold_left=True):
        p = d.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_before = Pt(4)
        a = p.add_run(left)
        a.bold = bold_left
        if right:
            p.add_run("\t" + right)
        return p

    def bullet(text: str):
        p = d.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(1)
        p.add_run(text)

    if r.get("name"):
        para(r["name"], bold=True, size=18, align="center")
    if r.get("headline"):
        para(r["headline"], size=11, align="center", color=(0x44, 0x44, 0x44))
    for c in r.get("contact") or []:
        para(c, size=9.5, align="center")

    if r.get("summary"):
        heading("Summary")
        para(r["summary"])
    if r.get("skills"):
        heading("Technical Skills")
        for s in r["skills"]:
            p = d.add_paragraph()
            if s.get("category"):
                p.add_run(f"{s['category']}: ").bold = True
            p.add_run(", ".join(s.get("items") or []))
    if r.get("experience"):
        heading("Experience")
        for e in r["experience"]:
            left_right(" | ".join(x for x in [e.get("company"), e.get("location")] if x),
                       e.get("dates") or "")
            if e.get("title"):
                para(e["title"], italic=True)
            for b in e.get("bullets") or []:
                bullet(b)
            for p in e.get("projects") or []:
                pp = d.add_paragraph()
                pp.paragraph_format.space_before = Pt(2)
                pp.add_run(p.get("name") or "").bold = True
                if p.get("tech"):
                    pp.add_run(f" | {p['tech']}")
                for b in p.get("bullets") or []:
                    bullet(b)
    if r.get("projects"):
        heading("Projects")
        for p in r["projects"]:
            pp = d.add_paragraph()
            pp.paragraph_format.space_before = Pt(3)
            pp.add_run(p.get("name") or "").bold = True
            if p.get("tech"):
                pp.add_run(f" | {p['tech']}")
            for b in p.get("bullets") or []:
                bullet(b)
    if r.get("education"):
        heading("Education")
        for line in r["education"]:
            para(line)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def cover_letter_docx(text: str, name: str = "") -> bytes:
    import docx
    from docx.shared import Inches, Pt
    d = docx.Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for block in re.split(r"\n\s*\n", text.strip()):
        p = d.add_paragraph(block.strip())
        p.paragraph_format.space_after = Pt(10)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def safe_filename(*parts: Any) -> str:
    s = "_".join(str(p) for p in parts if p)
    s = re.sub(r"[^A-Za-z0-9_\-]+", "_", s).strip("_")
    return s[:80] or "resume"
